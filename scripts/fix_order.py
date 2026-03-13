"""
HTML 문서 순서 기준으로 xlsx의 빨간(신규) 행 위치를 재정렬합니다.
- HTML에서 text-indent 패턴으로 엔드포인트 순서 추출
- 각 NEW 행을 HTML 순서상 직전 기존 행 바로 다음에 삽입
"""

import re, os, sys
import openpyxl
from openpyxl.styles import PatternFill
from copy import copy

ROOT  = r'C:\MIDAS\code\apiverification'
HTML  = os.path.join(ROOT, 'api_docs_zendesk', 'main_kr.html')
XLSX  = os.path.join(ROOT, 'manager_tasks_2026-03-12.xlsx')

LIGHT_RED = PatternFill(fill_type='solid', fgColor='FFFFE0E0')

def is_new_row(row):
    fill = row[0].fill
    return (fill and fill.patternType == 'solid' and fill.fgColor
            and fill.fgColor.type == 'rgb'
            and fill.fgColor.rgb.upper().endswith('FFE0E0'))

# ── HTML에서 엔드포인트 순서 추출 ──────────────────────────────
def parse_html_order():
    with open(HTML, encoding='utf-8') as f:
        content = f.read()
    pat = re.compile(
        r'text-indent:\s*10px[^>]*>\s*(/[A-Za-z0-9_\-/{}]+)\s*</p>',
        re.IGNORECASE
    )
    results = []
    for m in pat.finditer(content):
        path = m.group(1).strip().upper()
        if re.match(r'^/(DB|OPE|POST|VIEW|DOC|INFO)/', path):
            results.append(path)
    return results

def get_cell_value(cell):
    return str(cell.value).strip() if cell.value is not None else ''

def main():
    print('📄 HTML 순서 파싱 중...')
    html_order = parse_html_order()
    print(f'  HTML 엔드포인트 순서: {len(html_order)}개')

    html_idx = {p.upper(): i for i, p in enumerate(html_order)}

    mvld_items = sorted([(i, p) for p, i in html_idx.items() if 'MVLD' in p])
    print('\n  MVLD HTML 순서:')
    for i, p in mvld_items[:12]:
        print(f'    [{i:4d}] {p}')

    print('\n📂 xlsx 읽는 중...')
    wb = openpyxl.load_workbook(XLSX)
    ws = wb['Tasks']

    header_row = None
    for row in ws.iter_rows():
        if row[0].value == 'Order':
            header_row = row[0].row
            break
    print(f'  헤더 행: {header_row}')

    # ── 모든 데이터 행 읽기 ──────────────────────────────────────
    all_rows = []
    for row in ws.iter_rows(min_row=header_row + 1):
        ep  = get_cell_value(row[11])
        if not ep and row[0].value is None:
            continue
        tag = get_cell_value(row[2])
        new = is_new_row(row)
        data   = [ws.cell(row=row[0].row, column=c).value for c in range(1, ws.max_column + 1)]
        styles = []
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=row[0].row, column=c)
            styles.append({
                'fill': copy(cell.fill) if cell.has_style else None,
                'font': copy(cell.font) if cell.has_style else None,
                'border': copy(cell.border) if cell.has_style else None,
                'alignment': copy(cell.alignment) if cell.has_style else None,
                'number_format': cell.number_format,
            })
        all_rows.append({'ep': ep, 'tag': tag, 'is_new': new, 'data': data, 'styles': styles})

    old_rows = [r for r in all_rows if not r['is_new']]
    new_rows = [r for r in all_rows if r['is_new']]
    print(f'  전체 행: {len(all_rows)}  (OLD: {len(old_rows)}, NEW: {len(new_rows)})')

    # orig_idx 부여 (xlsx 행 순서 기록)
    for i, r in enumerate(old_rows):
        r['orig_idx'] = i

    # ── HTML rank 부여 ────────────────────────────────────────────
    def html_rank(ep):
        key = ep.upper()
        if key in html_idx:
            return (html_idx[key], key)

        seg1 = key.split('/')[1] if '/' in key else ''
        code = key.split('/')[-1] if '/' in key else key

        best_candidates = None

        # 우선 1: 코드에 '-' 가 있으면 첫 세그먼트(LCOM-CFSTEEL → LCOM) 기준으로 그룹 전체 검색
        if '-' in code:
            base = code.split('-')[0]          # e.g. 'LCOM'
            # '/SEG1/BASE-' 로 시작하는 HTML 경로 (가장 정확한 그룹)
            pfx_dash = '/' + seg1 + '/' + base + '-'
            cands = [(html_idx[k], k) for k in html_idx if k.startswith(pfx_dash)]
            if not cands:
                # '-' 없이 '/SEG1/BASE' 시작 (LCOM, MVHC 같은 경우)
                pfx_base = '/' + seg1 + '/' + base
                cands = [(html_idx[k], k) for k in html_idx if k.startswith(pfx_base)]
            if cands:
                best_candidates = cands

        # 우선 2: 길이를 줄여가며 접두사 매칭 (최소 4자, '-' 코드 이미 처리됨)
        if best_candidates is None:
            for prefix_len in range(len(code) - 1, 3, -1):
                prefix = '/' + seg1 + '/' + code[:prefix_len]
                cands = [(html_idx[k], k) for k in html_idx if k.startswith(prefix)]
                if cands:
                    best_candidates = cands
                    break

        # 폴백: 같은 seg1 내 알파벳 인접 3개
        if best_candidates is None:
            same = sorted([(html_idx[k], k) for k in html_idx
                           if '/' in k and k.split('/')[1] == seg1])
            if same:
                prev3 = [(i, k) for i, k in same if k <= key][-3:]
                nxt3  = [(i, k) for i, k in same if k > key][:3]
                best_candidates = prev3 + nxt3

        if not best_candidates:
            return (999999, key)

        best_candidates.sort()
        prev = [(i, k) for i, k in best_candidates if k <= key]
        if prev:
            last_i, _ = max(prev, key=lambda x: x[0])
            return (last_i + 0.5, key)
        else:
            return (best_candidates[0][0] - 0.5, key)

    for r in old_rows:
        r['rank'] = html_rank(r['ep']) if r['ep'] else (-1, '')
    for r in new_rows:
        r['rank'] = html_rank(r['ep'])

    # ── 전역 가장 가까운 OLD 선행자 찾기 ─────────────────────────
    # 각 NEW 행에 대해: rank <= new.rank 이면서 rank가 최대인 OLD 행 찾기
    # rank 동점이면 orig_idx(xlsx 후반)가 큰 쪽 선택
    from collections import defaultdict
    insertion_after = defaultdict(list)  # key = old 행의 orig_idx
    orphans = []  # OLD 선행자 없는 NEW 행

    for nr in new_rows:
        nr_rank = nr['rank']
        best_old = None
        best_old_rank = (float('-inf'), '')
        for or_ in old_rows:
            or_rank = or_['rank']
            if or_rank <= nr_rank:
                if or_rank > best_old_rank:
                    best_old_rank = or_rank
                    best_old = or_
                elif or_rank == best_old_rank and best_old is not None:
                    # 동점: xlsx에서 늦게 나오는 것 선택 (가장 가까운 위치)
                    if or_['orig_idx'] > best_old['orig_idx']:
                        best_old = or_
        if best_old is None:
            orphans.append(nr)
        else:
            insertion_after[best_old['orig_idx']].append(nr)

    print(f'  orphan NEW 행: {len(orphans)}개')

    # ── 최종 순서 구성 ─────────────────────────────────────────────
    final_rows = []
    # 선행 OLD 없는 NEW 행은 맨 앞에
    for nr in sorted(orphans, key=lambda r: r['rank']):
        final_rows.append(nr)

    for old_r in old_rows:
        final_rows.append(old_r)
        # 이 OLD 행 다음에 삽입할 NEW 행들 (rank 순)
        following = sorted(insertion_after.get(old_r['orig_idx'], []),
                           key=lambda r: r['rank'])
        final_rows.extend(following)

    print(f'정렬 후 총 행: {len(final_rows)}')

    # ── xlsx 재작성 ───────────────────────────────────────────────
    print('✍️  xlsx 재작성 중...')
    max_row = ws.max_row
    for r in range(max_row, header_row, -1):
        ws.delete_rows(r)

    for i, row_data in enumerate(final_rows):
        target_row = header_row + 1 + i
        ws.insert_rows(target_row)
        for col_idx, val in enumerate(row_data['data'], start=1):
            cell = ws.cell(row=target_row, column=col_idx)
            cell.value = val
            s = row_data['styles'][col_idx - 1]
            if s['fill']:      cell.fill      = s['fill']
            if s['font']:      cell.font      = s['font']
            if s['border']:    cell.border    = s['border']
            if s['alignment']: cell.alignment = s['alignment']
            cell.number_format = s['number_format']

    wb.save(XLSX)
    print(f'💾 저장: {XLSX}')

    # ── 검증 ──────────────────────────────────────────────────────
    print('\n📋 검증 - MVLD 관련 행:')
    wb2 = openpyxl.load_workbook(XLSX)
    ws2 = wb2['Tasks']
    for row in ws2.iter_rows(min_row=3):
        ep = row[11].value
        if ep and 'MVLD' in str(ep).upper() and 'INFO' not in str(ep).upper():
            new = is_new_row(row)
            mark = ' [NEW🔴]' if new else ''
            print(f'  Row{row[0].row:4d}  EP:{ep}{mark}')

if __name__ == '__main__':
    main()
