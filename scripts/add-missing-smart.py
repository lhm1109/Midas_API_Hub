"""
missing-endpoints → manager_tasks xlsx 에 적절한 위치에 행 삽입
- /{id} 항목 제외
- 이전에 하단 추가된 행(Order > 900) 정리
- 각 경로를 같은 seg1 그룹 내 알파벳 순 위치에 삽입
- 삽입된 행: 옅은 빨간색 배경 (#FFE0E0)
"""

import csv, io, os, sys
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

ROOT   = os.path.join(os.path.dirname(__file__), '..')
CSV    = os.path.join(ROOT, 'missing-endpoints.csv')
XLSX   = os.path.join(ROOT, 'manager_tasks_2026-03-12.xlsx')

LIGHT_RED = PatternFill(fill_type='solid', fgColor='FFFFE0E0')  # ARGB: 알파FF + 옅은빨강

# ── CSV 로드 ──
def load_missing():
    entries = []
    with open(CSV, encoding='utf-8') as f:
        reader = csv.reader(f)
        header = next(reader)
        for row in reader:
            if len(row) < 2: continue
            path = row[1].strip()
            if not path: continue
            if path.lower().endswith('/{id}'): continue  # /{id} 제외
            entries.append({
                'tag':     row[0].strip() if len(row) > 0 else '',
                'path':    path,
                'mode':    row[2].strip() if len(row) > 2 else '',
                'methods': row[3].strip() if len(row) > 3 else '',
                'desc':    row[4].strip() if len(row) > 4 else '',
            })
    return entries

def split_path(p):
    parts = p.lstrip('/').split('/')
    parts += ['', '', '', '']
    return parts[0], parts[1], parts[2], parts[3]

# ── 삽입 위치 결정 ──
def find_insert_after(missing_path, missing_tag, existing):
    """
    existing: list of (path, row_idx, tag)  -- row_idx는 1-based
    핵심: 같은 seg1 그룹 내에서 '알파벳으로 가장 가까운 선행 경로'의 행에 삽입.
    즉, path 값이 가장 크면서 missing_path 이하인 항목 → 그 항목의 행 바로 다음.
    반환값: insert_after_row_idx (1-based)
    """
    seg1 = missing_path.split('/')[1].upper() if '/' in missing_path else ''
    missing_upper = missing_path.upper()

    # 1차: 같은 seg1 그룹
    group = [(p, r, t) for p, r, t in existing
             if '/' in p and p.split('/')[1].upper() == seg1]

    # 2차: 같은 Tag
    if not group:
        group = [(p, r, t) for p, r, t in existing if t == missing_tag]

    if not group:
        return None  # 맨 끝에 추가

    # missing_path보다 알파벳 작거나 같은 후보
    candidates = [(p, r) for p, r, t in group if p.upper() <= missing_upper]

    if not candidates:
        # missing_path가 그룹의 첫 항목보다 앞 → 첫 항목 바로 앞에 삽입
        group.sort(key=lambda x: x[1])
        return group[0][1] - 1

    # 후보 중 '알파벳 값이 가장 큰' 경로의 행 선택
    # (= missing_path 직전에 오는 가장 가까운 항목)
    _, best_row = max(candidates, key=lambda x: x[0].upper())
    return best_row

# ── 행 복사 (스타일 유지용) ──
def copy_row_style(ws, src_row, dest_row):
    """src_row의 셀 스타일을 dest_row에 복사"""
    for col in range(1, ws.max_column + 2):
        src_cell = ws.cell(row=src_row, column=col)
        dst_cell = ws.cell(row=dest_row, column=col)
        if src_cell.has_style:
            dst_cell.font      = copy(src_cell.font)
            dst_cell.border    = copy(src_cell.border)
            dst_cell.alignment = copy(src_cell.alignment)

# ── 행 이동 (아래로) ──
def shift_rows_down(ws, from_row, count):
    """from_row 이상의 모든 행을 count만큼 아래로 이동"""
    max_row = ws.max_row
    max_col = ws.max_column

    for row in range(max_row, from_row - 1, -1):
        for col in range(1, max_col + 1):
            src = ws.cell(row=row, column=col)
            dst = ws.cell(row=row + count, column=col)
            dst.value = src.value
            if src.has_style:
                dst.font      = copy(src.font)
                dst.border    = copy(src.border)
                dst.fill      = copy(src.fill)
                dst.number_format = src.number_format
                dst.alignment = copy(src.alignment)
        # 원본 행 지우기
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.value = None
            cell.fill  = PatternFill(fill_type=None)

# ── 메인 ──
def main():
    print('📂 xlsx 읽는 중...')
    wb = load_workbook(XLSX)
    ws = wb['Tasks']

    # ─ Step 1: 헤더 행 찾기 (Order 컬럼) ─
    header_row = None
    for r in ws.iter_rows():
        if r[0].value == 'Order':
            header_row = r[0].row
            break
    if not header_row:
        print('❌ 헤더 행 없음'); sys.exit(1)
    print(f'헤더 행: {header_row}')

    # ─ Step 2: 이전 추가 행 제거 (Order > 900 또는 빨간 fill) ─
    rows_to_delete = []
    for row in ws.iter_rows(min_row=header_row + 1):
        order_val = row[0].value
        fill_val  = row[0].fill
        is_red = (fill_val and fill_val.patternType == 'solid' and
                  fill_val.fgColor and fill_val.fgColor.type == 'rgb' and
                  fill_val.fgColor.rgb.upper().endswith('FFE0E0'))
        if (isinstance(order_val, (int, float)) and order_val > 900) or is_red:
            rows_to_delete.append(row[0].row)

    # 아래서부터 삭제
    for r in sorted(rows_to_delete, reverse=True):
        ws.delete_rows(r)
    print(f'🗑️  이전 추가 행 {len(rows_to_delete)}개 제거')

    # ─ Step 3: 기존 EndPoint 맵 구성 ─
    existing = []  # (path, row_idx, tag)
    for row in ws.iter_rows(min_row=header_row + 1):
        ep  = row[11].value  # col 12 = index 11
        tag = row[2].value   # col 3  = index 2
        if ep and str(ep).strip():
            existing.append((str(ep).strip(), row[0].row, str(tag).strip() if tag else ''))
    print(f'기존 EndPoint 행: {len(existing)}개')

    # ─ Step 4: missing 로드 및 삽입 계획 ─
    missing = load_missing()
    existing_set = {p.lower() for p, _, _ in existing}

    insert_plan = {}  # after_row_idx -> [entry, ...]
    skip = 0
    for entry in missing:
        if entry['path'].lower() in existing_set:
            skip += 1
            continue
        after = find_insert_after(entry['path'], entry['tag'], existing)
        key = after if after is not None else 'end'
        insert_plan.setdefault(key, []).append(entry)

    total_planned = sum(len(v) for v in insert_plan.values())
    print(f'스킵(기존): {skip}개 / 삽입 예정: {total_planned}개')

    # ─ Step 5: 하단→상단 순으로 삽입 ─
    sorted_keys = sorted(
        insert_plan.keys(),
        key=lambda k: ws.max_row if k == 'end' else k,
        reverse=True
    )

    total_added = 0
    for key in sorted_keys:
        entries = sorted(insert_plan[key], key=lambda e: e['path'])
        insert_at = (ws.max_row + 1) if key == 'end' else key + 1

        for i, entry in enumerate(entries):
            actual_row = insert_at + i
            ws.insert_rows(actual_row)

            seg1, seg2, seg3, seg4 = split_path(entry['path'])

            ws.cell(row=actual_row, column=1).value  = None           # Order
            ws.cell(row=actual_row, column=2).value  = None           # Product
            ws.cell(row=actual_row, column=3).value  = entry['tag']   # Tab
            ws.cell(row=actual_row, column=4).value  = entry['tag']   # Group
            ws.cell(row=actual_row, column=5).value  = entry['desc']  # sub1 = desc
            ws.cell(row=actual_row, column=8).value  = seg1           # seg1
            ws.cell(row=actual_row, column=9).value  = seg2           # seg2
            ws.cell(row=actual_row, column=10).value = seg3           # seg3
            ws.cell(row=actual_row, column=11).value = seg4           # seg4
            ws.cell(row=actual_row, column=12).value = entry['path']  # EndPoint
            ws.cell(row=actual_row, column=13).value = entry['mode']  # mode
            ws.cell(row=actual_row, column=22).value = entry['methods'] # remark

            # 빨간 배경
            for col in range(1, 25):
                ws.cell(row=actual_row, column=col).fill = LIGHT_RED

            total_added += 1

    print(f'\n✅ 총 {total_added}개 행 삽입 완료')

    wb.save(XLSX)
    print(f'💾 저장: {XLSX}')

    # ─ 검증: MVLD 근처 ─
    print('\n📋 검증 - MVLD 관련 행:')
    wb2 = load_workbook(XLSX)
    ws2 = wb2['Tasks']
    for row in ws2.iter_rows(min_row=3):
        ep = row[11].value
        if ep and 'MVLD' in str(ep).upper():
            fill = row[0].fill
            is_new = (fill and fill.fgColor and fill.fgColor.type == 'rgb'
                      and fill.fgColor.rgb in ('FFFFE0E0', 'FFE0E0'))
            marker = ' [NEW🔴]' if is_new else ''
            print(f'  Row{row[0].row}  EP:{ep}{marker}')

if __name__ == '__main__':
    main()
